"""네이버 증권(Npay 증권)에서 코스피200 지수 데이터를 수집하는 스크립트.

수집 데이터
  1) 현재(또는 장 마감) 지수: 현재가, 전일비, 등락률, 시가/고가/저가, 거래량, 거래대금
  2) 일별 시세: 날짜별 종가, 전일비, 등락률, 시가, 고가, 저가
  3) 시가총액 상위 200개 종목 (페이징 수집): 순위, 종목코드, 종목명, 현재가, 전일비, 등락률, 거래량, 거래대금, 시가총액

시가총액 상위 200개에 대해
  - 코스피200 '공식 구성종목'을 무료로 제공하는 곳이 없어(네이버는 종료, KRX 는 로그인 필요) 대신 네이버의
    코스피 시가총액 순위를 페이지 단위(1회 최대 100개)로 넘겨 가며 ETF/ETN 을 뺀 상위 200개 종목을 수집합니다.
  - 공식 코스피200 구성종목과는 일부 다를 수 있으므로 '시가총액 상위 200' 으로 이해하세요.

왜 BeautifulSoup 이 아니라 JSON 을 쓰나요?
  - https://stock.naver.com/market/stock/kr 은 자바스크립트로 화면을 그리는 페이지라, 받아온 HTML 에는
    지수 숫자가 들어 있지 않습니다. (표에는 머리글만 있어 BeautifulSoup 으로는 읽을 수 없음)
  - 예전에 HTML 표로 코스피200을 제공하던 finance.naver.com 페이지는 서비스가 종료되었습니다 (410 Gone).
  - 대신 이 페이지가 화면을 그릴 때 호출하는 JSON 주소를 requests 로 직접 호출합니다.
    (JSON 은 HTML 이 아니므로 파싱에 BeautifulSoup 이 필요 없습니다.)

필요 패키지: pip install requests            (엑셀 저장 시: pip install openpyxl)

사용 예:
    python kospi200_crawler.py                        # 현재 지수 + 최근 10일 시세 출력
    python kospi200_crawler.py --days 60 --out kospi200.xlsx
    python kospi200_crawler.py --days 30 --out kospi200.csv   # 시세는 kospi200.csv, 종목은 kospi200_stocks.csv
    python kospi200_crawler.py --stocks 0             # 종목 수집 생략 (지수만)

주의: 네이버가 공개 문서로 제공하는 API 가 아니라 웹 페이지가 내부적으로 쓰는 주소이므로,
      예고 없이 바뀔 수 있습니다. 개인 학습 용도로 과도한 반복 호출은 피하세요.
"""

import argparse
import csv
import json
import sys
import time

import requests

REALTIME_URL = "https://polling.finance.naver.com/api/realtime/domestic/index/KPI200"
DAILY_URL = "https://m.stock.naver.com/api/index/KPI200/price"
PAGE_SIZE = 60
STOCKS_URL = "https://m.stock.naver.com/api/stocks/marketValue/KOSPI"  # 코스피 시가총액 순위 (페이징)
STOCK_PAGE_SIZE = 100  # 서버 최대값 (101 이상이면 400 오류)

HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
        "(KHTML, like Gecko) Chrome/124.0 Safari/537.36"
    ),
    "Referer": "https://stock.naver.com/",
    "Accept": "application/json",
}

DAILY_FIELDS = ["날짜", "종가", "전일비", "등락률(%)", "시가", "고가", "저가"]
STOCK_FIELDS = ["순위", "종목코드", "종목명", "현재가", "전일비", "등락률(%)",
                "거래량(주)", "거래대금(원)", "시가총액(원)"]


def to_float(text):
    """'1,112.16' → 1112.16, '-0.17' → -0.17. 변환할 수 없으면 None."""
    try:
        return float(str(text).replace(",", ""))
    except (TypeError, ValueError):
        return None


def get_json(session, url, params=None, retries=2):
    for attempt in range(retries + 1):
        try:
            resp = session.get(url, params=params, headers=HEADERS, timeout=10)
            resp.raise_for_status()
            return resp.json()
        except (requests.RequestException, ValueError) as exc:
            if attempt == retries:
                raise RuntimeError(f"요청 실패: {url} ({exc})") from exc
            time.sleep(1.5 * (attempt + 1))


def get_realtime(session):
    """코스피200 현재(장 마감 후에는 종가) 지수 정보를 dict 로 돌려준다."""
    data = get_json(session, REALTIME_URL)["datas"][0]
    return {
        "지수명": data["stockName"],
        "기준시각": data["localTradedAt"],
        "장상태": data["marketStatus"],          # OPEN / CLOSE
        "현재가": to_float(data["closePrice"]),
        "전일비": to_float(data["compareToPreviousClosePrice"]),
        "등락률(%)": to_float(data["fluctuationsRatio"]),
        "시가": to_float(data["openPrice"]),
        "고가": to_float(data["highPrice"]),
        "저가": to_float(data["lowPrice"]),
        "거래량(주)": to_float(data.get("accumulatedTradingVolumeRaw")),
        "거래대금(원)": to_float(data.get("accumulatedTradingValueRaw")),
    }


def get_daily(session, days):
    """최근 days 거래일의 일별 시세를 최신순 리스트로 돌려준다."""
    rows, page = [], 1
    while len(rows) < days:
        items = get_json(session, DAILY_URL, {"pageSize": PAGE_SIZE, "page": page})
        if not items:
            break
        for it in items:
            rows.append({
                "날짜": it["localTradedAt"],
                "종가": to_float(it["closePrice"]),
                "전일비": to_float(it["compareToPreviousClosePrice"]),
                "등락률(%)": to_float(it["fluctuationsRatio"]),
                "시가": to_float(it["openPrice"]),
                "고가": to_float(it["highPrice"]),
                "저가": to_float(it["lowPrice"]),
            })
        if len(items) < PAGE_SIZE:
            break
        page += 1
        time.sleep(0.5)
    return rows[:days]


def get_top_stocks(session, count=200, progress=None):
    """코스피 시가총액 상위 count 개 종목(ETF/ETN 제외)을 페이지 단위로 넘겨 가며 수집한다.

    한 번에 최대 100개만 받을 수 있고 ETF/ETN 이 순위에 섞여 있으므로,
    일반 종목이 count 개 모일 때까지(또는 마지막 페이지까지) 다음 페이지를 계속 요청한다.
    progress(page, collected) 를 주면 페이지를 하나 받을 때마다 호출한다.
    """
    stocks, seen, page = [], set(), 1
    while len(stocks) < count:
        data = get_json(session, STOCKS_URL, {"page": page, "pageSize": STOCK_PAGE_SIZE})
        items = data.get("stocks", [])
        if not items:
            break
        for it in items:
            code = it["itemCode"]
            # 페이지를 넘기는 동안 순위가 바뀌어 같은 종목이 두 번 나올 수 있으므로 중복 제거
            if it.get("stockEndType") != "stock" or code in seen:
                continue
            seen.add(code)
            stocks.append({
                "순위": len(stocks) + 1,
                "종목코드": code,
                "종목명": it["stockName"],
                "현재가": to_float(it["closePriceRaw"]),
                "전일비": to_float(it["compareToPreviousClosePriceRaw"]),
                "등락률(%)": to_float(it["fluctuationsRatio"]),
                "거래량(주)": to_float(it["accumulatedTradingVolumeRaw"]),
                "거래대금(원)": to_float(it["accumulatedTradingValueRaw"]),
                "시가총액(원)": to_float(it["marketValueRaw"]),
            })
            if len(stocks) == count:
                break
        if progress:
            progress(page, len(stocks))
        if page * STOCK_PAGE_SIZE >= data.get("totalCount", 0):
            break
        page += 1
        time.sleep(0.3)
    return stocks


def save_excel(realtime, daily, path, stocks=None):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError:
        raise RuntimeError("엑셀 저장에는 openpyxl 이 필요합니다: pip install openpyxl")

    wb = Workbook()

    def style_header(ws, widths):
        for col, width in enumerate(widths, 1):
            cell = ws.cell(row=1, column=col)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="2F5597")
            cell.alignment = Alignment(horizontal="center")
            ws.column_dimensions[cell.column_letter].width = width
        ws.freeze_panes = "A2"

    ws = wb.active
    ws.title = "일별시세"
    ws.append(DAILY_FIELDS)
    for row in daily:
        ws.append([row[f] for f in DAILY_FIELDS])
    for r in range(2, ws.max_row + 1):
        for c in range(2, 8):
            ws.cell(row=r, column=c).number_format = "#,##0.00"
    style_header(ws, [12, 12, 10, 11, 12, 12, 12])
    ws.auto_filter.ref = ws.dimensions

    ws2 = wb.create_sheet("현재지수")
    ws2.append(["항목", "값"])
    for key, value in realtime.items():
        ws2.append([key, value])
    for r in range(2, ws2.max_row + 1):
        if isinstance(ws2.cell(row=r, column=2).value, float):
            ws2.cell(row=r, column=2).number_format = "#,##0.00"
    style_header(ws2, [16, 28])

    if stocks:
        ws3 = wb.create_sheet(f"시총상위{len(stocks)}")
        ws3.append(STOCK_FIELDS)
        for row in stocks:
            ws3.append([row[f] for f in STOCK_FIELDS])
        for r in range(2, ws3.max_row + 1):
            ws3.cell(row=r, column=1).alignment = Alignment(horizontal="center")
            ws3.cell(row=r, column=2).alignment = Alignment(horizontal="center")
            for c in (4, 5, 7, 8, 9):
                ws3.cell(row=r, column=c).number_format = "#,##0"
            ws3.cell(row=r, column=6).number_format = "0.00"
        style_header(ws3, [6, 10, 24, 12, 10, 10, 16, 20, 22])
        ws3.auto_filter.ref = ws3.dimensions
    wb.save(path)


def write_csv(path, fields, rows):
    with open(path, "w", newline="", encoding="utf-8-sig") as f:  # utf-8-sig: 엑셀에서 한글 깨짐 방지
        writer = csv.DictWriter(f, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)


def save(realtime, daily, path, stocks=None):
    """확장자에 따라 저장한다. .csv 는 시세(path)와 종목(<이름>_stocks.csv) 두 파일로 나눠 저장한다."""
    lower = path.lower()
    if lower.endswith(".xlsx"):
        save_excel(realtime, daily, path, stocks)
    elif lower.endswith(".csv"):
        write_csv(path, DAILY_FIELDS, daily)
        if stocks:
            write_csv(path[:-4] + "_stocks.csv", STOCK_FIELDS, stocks)
    else:
        with open(path, "w", encoding="utf-8") as f:
            json.dump({"realtime": realtime, "daily": daily, "stocks": stocks or []},
                      f, ensure_ascii=False, indent=2)


def main():
    parser = argparse.ArgumentParser(description="코스피200 지수 데이터 크롤러 (네이버 증권)")
    parser.add_argument("--days", type=int, default=10, help="가져올 일별 시세 일수 (기본 10)")
    parser.add_argument("--stocks", type=int, default=200,
                        help="수집할 시가총액 상위 종목 수 (기본 200, 0 이면 생략)")
    parser.add_argument("--out", help="저장 파일 (.xlsx, .csv, .json). 생략하면 저장하지 않고 출력만 함")
    args = parser.parse_args()

    if hasattr(sys.stdout, "reconfigure"):  # 윈도우 콘솔에서 한글 출력 오류 방지
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")

    session = requests.Session()
    stocks = []
    try:
        realtime = get_realtime(session)
        daily = get_daily(session, args.days)
        if args.stocks > 0:
            print(f"시가총액 상위 {args.stocks}개 종목 수집 중 (ETF/ETN 제외, 페이지당 {STOCK_PAGE_SIZE}개)")
            stocks = get_top_stocks(
                session, args.stocks,
                progress=lambda page, n: print(f"  페이지 {page} 수집 → 종목 {n}개 확보"))
            print()
    except (RuntimeError, KeyError, IndexError) as exc:
        sys.exit(f"데이터를 가져오지 못했습니다: {exc}\n(네이버 API 주소나 형식이 바뀌었을 수 있습니다.)")

    status = "장중" if realtime["장상태"] == "OPEN" else "장 마감"
    print(f"[{realtime['지수명']}] {realtime['현재가']:,.2f}  "
          f"({realtime['전일비']:+,.2f}, {realtime['등락률(%)']:+.2f}%)  {status} · {realtime['기준시각']}")
    print(f"  시가 {realtime['시가']:,.2f}  고가 {realtime['고가']:,.2f}  저가 {realtime['저가']:,.2f}"
          f"  거래량 {realtime['거래량(주)']:,.0f}주\n")

    print(f"{'날짜':<12}{'종가':>10}{'전일비':>9}{'등락률':>9}{'시가':>10}{'고가':>10}{'저가':>10}")
    for r in daily:
        print(f"{r['날짜']:<12}{r['종가']:>10,.2f}{r['전일비']:>+9,.2f}{r['등락률(%)']:>+8.2f}%"
              f"{r['시가']:>10,.2f}{r['고가']:>10,.2f}{r['저가']:>10,.2f}")

    if stocks:
        print(f"\n시가총액 상위 {len(stocks)}개 종목")
        print(f"{'순위':>4}  {'종목코드':<8}{'종목명':<20}{'현재가':>11}{'등락률':>9}{'시가총액(조원)':>16}")
        for s in stocks:
            print(f"{s['순위']:>4}  {s['종목코드']:<8}{s['종목명']:<20}{s['현재가']:>11,.0f}"
                  f"{s['등락률(%)']:>+8.2f}%{s['시가총액(원)'] / 1e12:>16,.2f}")

    if args.out:
        try:
            save(realtime, daily, args.out, stocks)
        except (OSError, RuntimeError) as exc:
            sys.exit(f"저장 실패: {exc}")
        print(f"\n저장 완료: {args.out}")


if __name__ == "__main__":
    main()
