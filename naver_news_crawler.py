"""네이버 검색 결과(뉴스)에서 신문기사를 크롤링하는 스크립트.

1) 네이버 검색 결과 페이지에서 기사 목록(제목, 언론사, 시간, 요약, 링크)을 수집하고
2) 각 기사 링크에 접속해 본문을 추출한 뒤
3) 화면에 출력하고 JSON 또는 CSV 파일로 저장합니다.

필요 패키지: pip install requests beautifulsoup4

사용 예:
    python naver_news_crawler.py
    python naver_news_crawler.py --limit 5 --out result.csv
    python naver_news_crawler.py --out result.xlsx        # 엑셀로 저장 (openpyxl 필요)
    python naver_news_crawler.py --no-body                # 목록만 수집 (빠름)
    python naver_news_crawler.py --url "https://search.naver.com/search.naver?where=nexearch&query=AI"

주의:
- 네이버 페이지 구조(HTML)는 자주 바뀔 수 있습니다. 목록이 비어 나오면 parse_search() 의 선택자를 점검하세요.
- 각 언론사 사이트의 이용약관/robots.txt 를 확인하고, 개인 학습 용도로 서버에 부담을 주지 않는 범위에서만 사용하세요.
  (기본적으로 기사 요청 사이에 --delay 초만큼 쉽니다.)
- 수집한 기사의 저작권은 각 언론사에 있습니다.
"""

import argparse
import csv
import json
import re
import sys
import time
from urllib.parse import parse_qs, urljoin, urlparse

import requests
from bs4 import BeautifulSoup

DEFAULT_URL = (
    "https://search.naver.com/search.naver?where=nexearch&sm=top_hty&fbm=0"
    "&ie=utf8&query=%EB%8F%99%EB%AC%BC%EC%9A%A9%EC%9D%98%EC%95%BD%ED%92%88&ackey=5klo6jet"
)

HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
        "(KHTML, like Gecko) Chrome/124.0 Safari/537.36"
    ),
    "Accept-Language": "ko-KR,ko;q=0.9,en;q=0.5",
}

# 언론사별로 구조가 달라서, 자주 쓰이는 본문 선택자를 순서대로 시도합니다.
BODY_SELECTORS = [
    "#dic_area",                    # 네이버 뉴스
    "#newsct_article",
    "#articleBodyContents",
    "#articleBody",
    "#article-view-content-div",
    "#newsContent",
    "#article_body",
    "[itemprop='articleBody']",
    ".article_body",
    ".article-body",
    ".news_body",
    ".view_con",
    "article",
]
# 본문에서 제거할 요소 (광고, 스크립트, 관련기사 등)
NOISE_SELECTORS = [
    "script", "style", "noscript", "iframe", "figure", "figcaption",
    "aside", "button", "nav", "header", "footer",  # form 은 제외: 페이지 전체를 <form> 으로 감싼 사이트가 있음
    ".ad", ".ads", ".advertisement", ".related", ".relation",
    ".reporter", ".byline", ".copyright", ".sns", ".share",
]
MIN_BODY_LEN = 200  # 이보다 짧으면 본문이 아니라고 보고 다음 방법 시도


def clean_text(text):
    """공백/개행을 정리한다."""
    text = text.replace("\xa0", " ")
    lines = [re.sub(r"[ \t]+", " ", ln).strip() for ln in text.splitlines()]
    return "\n".join(ln for ln in lines if ln)


def one_line(node):
    """태그 안의 텍스트를 한 줄로 합친다 (<mark> 등으로 쪼개진 단어가 붙어 있도록 strip 하지 않음)."""
    text = node.get_text() if node else ""
    return re.sub(r"\s+", " ", text.replace("새 창 열림", "")).strip()


META_REFRESH = re.compile(
    r"""<meta[^>]+http-equiv=["']refresh["'][^>]+content=["'][^"']*url=([^"']+)["']""", re.I)


def fetch(session, url, timeout=10, retries=2, follow_refresh=True):
    """URL 을 가져와 HTML 문자열로 돌려준다. 실패하면 None."""
    for attempt in range(retries + 1):
        try:
            resp = session.get(url, headers=HEADERS, timeout=timeout)
            resp.raise_for_status()
            # 일부 언론사는 EUC-KR 인데 헤더에 charset 이 없어 깨질 수 있음
            if resp.encoding is None or resp.encoding.lower() == "iso-8859-1":
                resp.encoding = resp.apparent_encoding
            # <meta http-equiv="refresh"> 로 다른 주소로 넘기는 사이트(예: 동아사이언스) 처리 (1회만)
            m = META_REFRESH.search(resp.text) if follow_refresh else None
            if m:
                return fetch(session, urljoin(resp.url, m.group(1)), timeout, retries, follow_refresh=False)
            return resp.text
        except requests.RequestException as exc:
            if attempt == retries:
                print(f"  [요청 실패] {url} ({exc})", file=sys.stderr)
                return None
            time.sleep(1.5 * (attempt + 1))


def parse_search(html):
    """검색 결과 HTML 에서 기사 목록을 추출한다."""
    soup = BeautifulSoup(html, "html.parser")
    articles, seen = [], set()

    # 클래스명은 난독화되어 매번 바뀌므로, 안정적인 data-heatmap-target 속성을 기준으로 잡는다.
    for title_a in soup.select("a[data-heatmap-target='.tit']"):
        href = title_a.get("href", "")
        if not href.startswith("http") or href in seen:
            continue
        seen.add(href)

        # 제목 링크에서 위로 올라가며 언론사명이 들어 있는 가장 가까운 묶음을 찾는다.
        item = title_a
        while item.parent is not None and not item.select_one(".sds-comps-profile-info-title-text"):
            item = item.parent

        press = one_line(item.select_one(".sds-comps-profile-info-title-text"))
        posted = one_line(item.select_one(".sds-comps-profile-info-subtext"))
        summary_a = item.select_one("a[data-heatmap-target='.body']")

        articles.append({
            "title": one_line(title_a),
            "press": press,
            "posted": posted,
            "summary": one_line(summary_a),
            "url": href,
            "body": "",
        })
    return articles


def extract_fusion(html):
    """조선일보/조선비즈 등 Arc Fusion 기반 사이트: 본문이 HTML 이 아니라 자바스크립트 JSON 안에 들어 있다."""
    marker = "Fusion.globalContent="
    start = html.find(marker)
    if start < 0:
        return ""
    try:
        data, _ = json.JSONDecoder().raw_decode(html[start + len(marker):])
    except ValueError:
        return ""
    paragraphs = [
        BeautifulSoup(el.get("content", ""), "html.parser").get_text(" ", strip=True)
        for el in data.get("content_elements", [])
        if el.get("type") == "text"
    ]
    return clean_text("\n".join(p for p in paragraphs if p))


def extract_body(html):
    """기사 페이지 HTML 에서 본문 텍스트를 추출한다."""
    text = extract_fusion(html)
    if len(text) >= MIN_BODY_LEN:
        return text

    soup = BeautifulSoup(html, "html.parser")
    for sel in NOISE_SELECTORS:
        for tag in soup.select(sel):
            tag.decompose()

    # 1) 알려진 본문 선택자
    for sel in BODY_SELECTORS:
        node = soup.select_one(sel)
        if node:
            text = clean_text(node.get_text("\n"))
            if len(text) >= MIN_BODY_LEN:
                return text

    # 2) 범용 방법: <p> 텍스트가 가장 많이 모여 있는 부모를 본문으로 간주
    best, best_len = None, 0
    for parent in {p.parent for p in soup.find_all("p") if p.parent}:
        total = sum(len(p.get_text(strip=True)) for p in parent.find_all("p", recursive=False))
        if total > best_len:
            best, best_len = parent, total
    if best is not None and best_len >= MIN_BODY_LEN:
        paragraphs = (p.get_text(" ", strip=True) for p in best.find_all("p", recursive=False))
        return clean_text("\n".join(paragraphs))

    # 3) 마지막 수단: 메타 태그의 요약문
    meta = soup.find("meta", property="og:description") or soup.find("meta", attrs={"name": "description"})
    return clean_text(meta["content"]) if meta and meta.get("content") else ""


def save_excel(articles, path):
    """기사 목록을 서식이 적용된 엑셀(.xlsx) 파일로 저장한다. (openpyxl 필요)"""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError:
        raise RuntimeError("엑셀 저장에는 openpyxl 이 필요합니다: pip install openpyxl")

    illegal = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")  # 엑셀이 허용하지 않는 제어문자

    def cell_text(value):
        return illegal.sub("", value or "")[:32000]  # 셀 하나의 최대 길이는 32,767자

    columns = [  # (제목, 키, 너비)
        ("번호", None, 6),
        ("제목", "title", 45),
        ("언론사", "press", 12),
        ("시간", "posted", 12),
        ("요약", "summary", 40),
        ("링크", "url", 40),
        ("본문", "body", 80),
    ]
    wb = Workbook()
    ws = wb.active
    ws.title = "뉴스"
    ws.append([name for name, _, _ in columns])
    for col, (_, _, width) in enumerate(columns, 1):
        cell = ws.cell(row=1, column=col)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2F5597")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[cell.column_letter].width = width

    for no, art in enumerate(articles, 1):
        ws.append([no] + [cell_text(art.get(key)) for _, key, _ in columns[1:]])
        row = ws.max_row
        ws.cell(row=row, column=1).alignment = Alignment(horizontal="center", vertical="top")
        for col in range(2, len(columns) + 1):
            ws.cell(row=row, column=col).alignment = Alignment(wrap_text=True, vertical="top")
        link = ws.cell(row=row, column=6)
        if art.get("url"):
            link.hyperlink = art["url"]
            link.font = Font(color="0563C1", underline="single")
        # 본문이 매우 길어도 행이 화면을 다 차지하지 않도록 높이를 제한
        ws.row_dimensions[row].height = 90

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(path)


def save(articles, path):
    if path.lower().endswith(".xlsx"):
        save_excel(articles, path)
    elif path.lower().endswith(".csv"):
        with open(path, "w", newline="", encoding="utf-8-sig") as f:  # utf-8-sig: 엑셀에서 한글 깨짐 방지
            writer = csv.DictWriter(f, fieldnames=["title", "press", "posted", "summary", "url", "body"])
            writer.writeheader()
            writer.writerows(articles)
    else:
        with open(path, "w", encoding="utf-8") as f:
            json.dump(articles, f, ensure_ascii=False, indent=2)


def main():
    parser = argparse.ArgumentParser(description="네이버 검색 결과의 뉴스 기사 크롤러")
    parser.add_argument("--url", default=DEFAULT_URL, help="네이버 검색 결과 URL (기본: '동물용의약품' 검색)")
    parser.add_argument("--limit", type=int, default=10, help="가져올 기사 수 (기본 10)")
    parser.add_argument("--delay", type=float, default=1.0, help="기사 요청 사이 대기 시간(초, 기본 1.0)")
    parser.add_argument("--no-body", action="store_true", help="본문은 가져오지 않고 목록만 수집")
    parser.add_argument("--out", default="naver_news.json", help="저장 파일 (.json, .csv, .xlsx)")
    args = parser.parse_args()

    if hasattr(sys.stdout, "reconfigure"):  # 윈도우 콘솔에서 한글 출력 오류 방지
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")

    query = parse_qs(urlparse(args.url).query).get("query", ["?"])[0]
    print(f"검색어: {query}")

    session = requests.Session()
    html = fetch(session, args.url)
    if not html:
        sys.exit("검색 결과 페이지를 가져오지 못했습니다.")

    articles = parse_search(html)[: args.limit]
    if not articles:
        sys.exit("기사를 찾지 못했습니다. 네이버 페이지 구조가 바뀌었을 수 있습니다 (parse_search 선택자 점검).")
    print(f"기사 {len(articles)}건을 찾았습니다.\n")

    for i, art in enumerate(articles, 1):
        print(f"[{i}] {art['title']}")
        print(f"    {art['press']} · {art['posted']}")
        print(f"    {art['url']}")

        if not args.no_body:
            time.sleep(args.delay)
            page = fetch(session, art["url"])
            if page:
                art["body"] = extract_body(page)
        preview = art["body"] or art["summary"]
        if preview:
            shown = preview.replace("\n", " ")
            print(f"    내용: {shown[:150]}{'…' if len(shown) > 150 else ''}")
        print()

    save(articles, args.out)
    print(f"저장 완료: {args.out}")


if __name__ == "__main__":
    main()
